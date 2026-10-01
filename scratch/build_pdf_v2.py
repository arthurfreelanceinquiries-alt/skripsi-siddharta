"""Build PDF v2 parity DOCX v2 - body walk, APA tables, fig, equations, roman->arab footer."""
import re, html
from pathlib import Path
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, Table, TableStyle, PageBreak, Image, NextPageTemplate, Flowable
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

DOCX_PATH = r"Z:\SKRIPSII\SKRIPSI ARTHUR\skripsi siddharta\01_Naskah_Utama\Proposal_Skripsi_MLBB_GenZ_v2.docx"
PDF_PATH = r"Z:\SKRIPSII\SKRIPSI ARTHUR\skripsi siddharta\01_Naskah_Utama\Proposal_Skripsi_MLBB_GenZ_v2.pdf"
FIG_PATH = r"C:\Users\Arthur Reezan\AppData\Local\Temp\opencode\fig_v2.png"
if not Path(FIG_PATH).exists():
    FIG_PATH = r"Z:\SKRIPSII\SKRIPSI ARTHUR\skripsi siddharta\scratch\fig_v2.png"

pdfmetrics.registerFont(TTFont('TNR', r'C:\Windows\Fonts\times.ttf'))
pdfmetrics.registerFont(TTFont('TNR-Bold', r'C:\Windows\Fonts\timesbd.ttf'))
pdfmetrics.registerFont(TTFont('TNR-Italic', r'C:\Windows\Fonts\timesi.ttf'))
pdfmetrics.registerFont(TTFont('TNR-BoldItalic', r'C:\Windows\Fonts\timesbi.ttf'))
pdfmetrics.registerFontFamily('TNR', normal='TNR', bold='TNR-Bold', italic='TNR-Italic', boldItalic='TNR-BoldItalic')

class TOCLineFlowable(Flowable):
    def __init__(self, title, page, fontName='TNR', fontSize=12, leftIndent=0, spaceBefore=2, spaceAfter=2, usableWidth=396.85):
        Flowable.__init__(self)
        self.title = title
        self.page = str(page)
        self.fontName = fontName
        self.fontSize = fontSize
        self.leftIndent = leftIndent
        self.spaceBefore = spaceBefore
        self.spaceAfter = spaceAfter
        self.usableWidth = usableWidth
        self.leading = 14.0
        self.width = usableWidth
        avail = usableWidth - leftIndent
        tw = pdfmetrics.stringWidth(title, fontName, fontSize)
        pw = pdfmetrics.stringWidth(self.page, fontName, fontSize)
        if tw + pw + 12 <= avail:
            self.lines = [title]
        else:
            words = title.split()
            l1 = []
            l2 = []
            for w in words:
                test_l1 = ' '.join(l1 + [w])
                if pdfmetrics.stringWidth(test_l1, fontName, fontSize) <= avail:
                    l1.append(w)
                else:
                    l2.append(w)
            self.lines = [' '.join(l1), ' '.join(l2)]
        self.height = len(self.lines) * self.leading + spaceBefore + spaceAfter

    def wrap(self, aW, aH):
        return self.usableWidth, self.height

    def draw(self):
        canvas = self.canv
        canvas.saveState()
        canvas.setFont(self.fontName, self.fontSize)
        canvas.setFillColor(colors.black)
        if len(self.lines) == 1:
            y = self.spaceAfter + 2
            tw = pdfmetrics.stringWidth(self.lines[0], self.fontName, self.fontSize)
            pw = pdfmetrics.stringWidth(self.page, self.fontName, self.fontSize)
            x_title = self.leftIndent
            x_page = self.usableWidth - pw
            gap_start = x_title + tw + 4
            gap_end = x_page - 4
            gap_width = gap_end - gap_start
            canvas.drawString(x_title, y, self.lines[0])
            if gap_width > 8:
                dot_w = pdfmetrics.stringWidth('.', 'TNR', self.fontSize)
                num_dots = int(gap_width / dot_w)
                dots_str = '.' * num_dots
                dots_x = gap_end - num_dots * dot_w
                canvas.setFont('TNR', self.fontSize)
                canvas.drawString(dots_x, y, dots_str)
            canvas.setFont(self.fontName, self.fontSize)
            canvas.drawRightString(self.usableWidth, y, self.page)
        else:
            y1 = self.spaceAfter + 2 + self.leading
            y2 = self.spaceAfter + 2
            canvas.drawString(self.leftIndent, y1, self.lines[0])
            tw2 = pdfmetrics.stringWidth(self.lines[1], self.fontName, self.fontSize)
            pw = pdfmetrics.stringWidth(self.page, self.fontName, self.fontSize)
            x_title2 = self.leftIndent
            x_page = self.usableWidth - pw
            gap_start = x_title2 + tw2 + 4
            gap_end = x_page - 4
            gap_width = gap_end - gap_start
            canvas.drawString(x_title2, y2, self.lines[1])
            if gap_width > 8:
                dot_w = pdfmetrics.stringWidth('.', 'TNR', self.fontSize)
                num_dots = int(gap_width / dot_w)
                dots_str = '.' * num_dots
                dots_x = gap_end - num_dots * dot_w
                canvas.setFont('TNR', self.fontSize)
                canvas.drawString(dots_x, y2, dots_str)
            canvas.setFont(self.fontName, self.fontSize)
            canvas.drawRightString(self.usableWidth, y2, self.page)
        canvas.restoreState()

def esc(t):
    return html.escape(t, quote=False).replace('\u00a0',' ')

URL_RE = re.compile(r'(https?://[^\s\)\]]+)')

d = Document(DOCX_PATH)
rels = {}
for rel in d.part.rels.values():
    if rel.reltype == RT.HYPERLINK:
        rels[rel.rId] = rel.target_ref

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

def para_to_markup(para):
    if '\t' in para.text:
        # TOC/LOT/LOF or equation? equations also have tab but contain (3.1) - handle separately outside
        # This generic handles TOC entries
        parts_txt = para.text.split('\t')
        title = parts_txt[0].strip()
        page = parts_txt[-1].strip() if len(parts_txt)>1 else ''
        dots = '.'*52
        # preserve bold for toc1?
        is_bold = any((r.bold for r in para.runs if r.text.strip()))
        if is_bold:
            return f'<b>{esc(title)}</b> {dots} {esc(page)}'
        else:
            return f'{esc(title)} {dots} {esc(page)}'
    parts = []
    for child in para._p.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'r':
            txts = [n.text or '' for n in child.findall(f'{{{W_NS}}}t')]
            txt = ''.join(txts)
            if not txt:
                if child.find(f'{{{W_NS}}}tab') is not None:
                    txt = ' '
                elif child.find(f'{{{W_NS}}}br') is not None:
                    parts.append('<br/>')
                    continue
                else:
                    continue
            rpr = child.find(f'{{{W_NS}}}rPr')
            b=False; it=False
            if rpr is not None:
                b_el = rpr.find(f'{{{W_NS}}}b')
                if b_el is not None:
                    if b_el.get(f'{{{W_NS}}}val','1') not in ('0','false','off'):
                        b=True
                i_el = rpr.find(f'{{{W_NS}}}i')
                if i_el is not None:
                    if i_el.get(f'{{{W_NS}}}val','1') not in ('0','false','off'):
                        it=True
            e = esc(txt).replace('\n','<br/>').replace('\t',' ')
            # URL regex for plain-text URLs outside hyperlink
            # do linking after bold/italic? URLs are plain, so link inside
            # check if URL in txt and not already hyperlink -> wrap
            # To avoid breaking <b><i>, do regex on escaped txt before wrapping, then wrap?
            # Simpler: if URL present, split and link
            if URL_RE.search(txt) and not b and not it:
                # plain URL run
                def repl(m):
                    url=m.group(1)
                    trail=''
                    while url and url[-1] in '.,;':
                        trail=url[-1]+trail
                        url=url[:-1]
                    return f'<a href="{html.escape(url, quote=True)}" color="black">{html.escape(url)}</a>{trail}'
                e2 = URL_RE.sub(repl, esc(txt)).replace('\n','<br/>')
                parts.append(e2)
                continue
            if b and it:
                e = f'<b><i>{e}</i></b>'
            elif b:
                e = f'<b>{e}</b>'
            elif it:
                e = f'<i>{e}</i>'
            parts.append(e)
        elif tag == 'hyperlink':
            rId = child.get(f'{{{R_NS}}}id')
            url = rels.get(rId,'')
            htxt = ''.join([n.text or '' for n in child.iter() if n.tag.endswith('}t')])
            if not htxt and url:
                htxt=url
            e = esc(htxt)
            # check bold/italic inside hyperlink? look at first rPr
            # For refs URLs, plain; keep plain link black
            if url:
                parts.append(f'<a href="{html.escape(url, quote=True)}" color="black">{e}</a>')
            else:
                parts.append(e)
        elif tag in ('proofErr','bookmarkStart','bookmarkEnd'):
            continue
    s=''.join(parts)
    if not s.strip() and para.text.strip():
        # fallback
        s=esc(para.text).replace('\n','<br/>')
        # link URLs
        s=URL_RE.sub(lambda m: f'<a href="{html.escape(m.group(1), quote=True)}" color="black">{html.escape(m.group(1))}</a>', s)
    return s if s else ' '

def int_to_roman(n):
    vals=[(1000,'M'),(900,'CM'),(500,'D'),(400,'CD'),(100,'C'),(90,'XC'),(50,'L'),(40,'XL'),(10,'X'),(9,'IX'),(5,'V'),(4,'IV'),(1,'I')]
    r=''
    for v,s in vals:
        while n>=v:
            r+=s; n-=v
    return r.lower() if r else ''

# styles
usable = A4[0] - 4*cm - 3*cm
body_style = ParagraphStyle('Body', fontName='TNR', fontSize=12, leading=18, alignment=TA_JUSTIFY, textColor=colors.black, spaceAfter=0, spaceBefore=0, firstLineIndent=1.25*cm)
body_hang = ParagraphStyle('BodyHang', parent=body_style, leftIndent=1.25*cm, firstLineIndent=-0.40*cm)
body_noindent = ParagraphStyle('BodyNoInd', parent=body_style, firstLineIndent=0)
center12 = ParagraphStyle('Center12', parent=body_style, alignment=TA_CENTER, firstLineIndent=0)
left12 = ParagraphStyle('Left12', parent=body_style, alignment=TA_LEFT, firstLineIndent=0)
h1_style = ParagraphStyle('H1', parent=body_style, alignment=TA_CENTER, fontName='TNR-Bold', spaceBefore=12, spaceAfter=12, firstLineIndent=0, keepWithNext=True)
h2_style = ParagraphStyle('H2', parent=body_style, alignment=TA_LEFT, fontName='TNR-Bold', spaceBefore=12, spaceAfter=6, firstLineIndent=0, keepWithNext=True)
h3_style = ParagraphStyle('H3', parent=body_style, alignment=TA_LEFT, fontName='TNR-Bold', spaceBefore=6, spaceAfter=3, firstLineIndent=0, keepWithNext=True)
ref_style = ParagraphStyle('Ref', fontName='TNR', fontSize=12, leading=13.8, alignment=TA_JUSTIFY, textColor=colors.black, leftIndent=36, firstLineIndent=-36, spaceAfter=0, spaceBefore=0)
toc_style = ParagraphStyle('TOC', parent=body_style, alignment=TA_LEFT, firstLineIndent=0, spaceAfter=0, spaceBefore=0)
toc_style_b = ParagraphStyle('TOCb', parent=toc_style, fontName='TNR-Bold')
caption_style = ParagraphStyle('Caption', fontName='TNR-Bold', fontSize=11, leading=13.2, alignment=TA_LEFT, textColor=colors.black, spaceBefore=12, spaceAfter=6, firstLineIndent=0)
caption_center = ParagraphStyle('CaptionC', parent=caption_style, alignment=TA_CENTER)
source_style = ParagraphStyle('Source', fontName='TNR-Italic', fontSize=10, leading=12, alignment=TA_LEFT, textColor=colors.black, spaceBefore=2, spaceAfter=12, firstLineIndent=0)
source_center = ParagraphStyle('SourceC', parent=source_style, alignment=TA_CENTER)
cover14 = ParagraphStyle('Cover14', fontName='TNR-Bold', fontSize=14, leading=16.8, alignment=TA_CENTER, textColor=colors.black, firstLineIndent=0)
cover12 = ParagraphStyle('Cover12', fontName='TNR', fontSize=12, leading=18, alignment=TA_CENTER, textColor=colors.black, firstLineIndent=0)
cell12 = ParagraphStyle('Cell12', fontName='TNR', fontSize=12, leading=14, alignment=TA_LEFT, textColor=colors.black)
cell10 = ParagraphStyle('Cell10', fontName='TNR', fontSize=10, leading=12, alignment=TA_LEFT, textColor=colors.black)
cellH12 = ParagraphStyle('CellH12', parent=cell12, fontName='TNR-Bold', alignment=TA_CENTER)
cellH10 = ParagraphStyle('CellH10', parent=cell10, fontName='TNR-Bold', alignment=TA_CENTER)
eq_style = ParagraphStyle('Eq', parent=body_style, alignment=TA_CENTER, firstLineIndent=0)
eqnum_style = ParagraphStyle('EqNum', parent=body_style, alignment=TA_RIGHT, firstLineIndent=0)

# footer counters
_main_counter = [1]
def cover_footer(canvas, doc):
    canvas.saveState()
    canvas.restoreState()
def front_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont('TNR-Bold', 10)
    canvas.setFillColor(colors.black)
    roman = int_to_roman(canvas.getPageNumber())
    txt = f'Universitas Kristen Krida Wacana | {roman}'
    canvas.drawCentredString(A4[0]/2, 36, txt)
    canvas.restoreState()
def main_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont('TNR-Bold', 10)
    canvas.setFillColor(colors.black)
    num = _main_counter[0]
    txt = f'Universitas Kristen Krida Wacana | {num}'
    # right aligned at right margin
    canvas.drawRightString(A4[0]-3*cm, 36, txt)
    _main_counter[0]+=1
    canvas.restoreState()

# doc template
doc_tpl = BaseDocTemplate(PDF_PATH, pagesize=A4, leftMargin=4*cm, rightMargin=3*cm, topMargin=3*cm, bottomMargin=3*cm, title='Proposal Skripsi MLBB GenZ v2', author='FEB UKRIDA')
frame = Frame(doc_tpl.leftMargin, doc_tpl.bottomMargin, doc_tpl.width, doc_tpl.height, id='normal')
doc_tpl.addPageTemplates([PageTemplate(id='Cover', frames=frame, onPage=cover_footer),
                          PageTemplate(id='Front', frames=frame, onPage=front_footer),
                          PageTemplate(id='Main', frames=frame, onPage=main_footer)])

# find ref start
ref_start=None
for i,para in enumerate(d.paragraphs):
    if para.text.startswith('Ajzen, I. (1991)'):
        ref_start=i
        break
print('ref_start', ref_start)

# map para/table els
# Build story via body walk
story=[]
# start on Cover (default first template is Cover)
sect_idx=0
# Track previous was PageBreak to avoid doubles
def last_is_break():
    return len(story)>0 and isinstance(story[-1], PageBreak)

# Need to handle cover spacings: paras 0-5 have specific before/after
cover_bef_aft = {0:(6,18),1:(12,24),2:(18,0),3:(0,24),4:(24,36),5:(48,0)}

for child in d.element.body.iterchildren():
    tag=child.tag.split('}')[-1]
    if tag=='p':
        # find para
        para=None; p_idx=None
        for _i,_p in enumerate(d.paragraphs):
            if _p._p is child:
                para=_p; p_idx=_i; break
        if para is None:
            continue
        idx=p_idx
        xml=para._p.xml
        has_sect = ('<w:sectPr' in xml)
        # sect break paras are empty with sectPr (006,046). They contain sectPr inside pPr.
        # Distinguish from final sectPr? final BODY-SECTPR is tag sectPr, not p. So here only 006,046.
        is_sect_para = has_sect and para.text.strip()=='' and ('<w:drawing' not in xml)
        has_pb = 'pageBreakBefore' in xml
        has_drawing = '<w:drawing' in xml
        txt=para.text
        # SECT handling
        if is_sect_para:
            if sect_idx==0:
                story.append(NextPageTemplate('Front'))
                story.append(PageBreak())
            elif sect_idx==1:
                story.append(NextPageTemplate('Main'))
                story.append(PageBreak())
            else:
                story.append(PageBreak())
            sect_idx+=1
            continue
        # TOC field para (009) empty with TOC field? contains TOC but empty text
        if 'TOC \\o' in xml or ('fldChar' in xml and txt.strip()==''):
            # skip field char para, add tiny spacer? skip to avoid extra blank
            continue
        # empty para handling
        if txt.strip()=='':
            if has_drawing:
                # image para 127
                # insert image 14cm width
                try:
                    img = Image(FIG_PATH, width=14*cm, height=5.93*cm)
                    img.hAlign='CENTER'
                    story.append(img)
                    story.append(Spacer(1,6))
                except Exception as e:
                    print('img fail', e)
                continue
            else:
                # buffer empty -> spacer (but avoid spacer right after PageBreak? still spacer creates blank line? Use small spacer)
                # For 038,042,045 etc, add Spacer
                story.append(Spacer(1,6))
                continue
        # pageBreakBefore handling
        if has_pb:
            if not last_is_break():
                # check if previous is NextPageTemplate? NextPageTemplate is not break, but PageBreak after it is. For sect+heading combos, avoid double.
                # If last two are NextPageTemplate+PageBreak, skip
                if len(story)>=2 and isinstance(story[-1], PageBreak):
                    pass
                else:
                    story.append(PageBreak())
        # Now content
        st_name=para.style.name
        # Equation special: center paras with (3.1)/(3.2) and tab
        if '\t' in txt and ('(3.1)' in txt or '(3.2)' in txt):
            # split equation and number
            parts_txt = txt.split('\t')
            eq_txt = parts_txt[0].strip()
            num_txt = parts_txt[-1].strip()
            # Build markup preserving italic for eq part: use para runs? First run is eq italic, second is num.
            # Simplify: eq italic, num plain
            eq_markup = f'<i>{esc(eq_txt)}</i>'
            num_markup = esc(num_txt)
            eq_para = Paragraph(eq_markup, eq_style)
            num_para = Paragraph(num_markup, eqnum_style)
            t = Table([[eq_para, num_para]], colWidths=[usable*0.85, usable*0.15])
            t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),2),('RIGHTPADDING',(0,0),(-1,-1),2),('TOPPADDING',(0,0),(-1,-1),2),('BOTTOMPADDING',(0,0),(-1,-1),2)]))
            story.append(t)
            story.append(Spacer(1,6))
            continue
        # TOC entries (contain tab, not equation)
        if '\t' in txt:
            parts_txt = txt.split('\t')
            entry_title = parts_txt[0].strip()
            entry_page = parts_txt[-1].strip() if len(parts_txt)>1 else ''
            is_b = any((r.bold for r in para.runs if r.text.strip()))
            
            if st_name=='toc 1' or is_b:
                f_name = 'TNR-Bold' if is_b else 'TNR'
                l_indent = 0
                sp_before = 4 if is_b else 2
                sp_after = 2 if is_b else 2
            elif st_name=='toc 2':
                f_name = 'TNR'
                l_indent = 14.17 # 0.5 cm
                sp_before = 1
                sp_after = 1
            elif st_name=='toc 3':
                f_name = 'TNR'
                l_indent = 28.35 # 1.0 cm
                sp_before = 1
                sp_after = 1
            else:
                f_name = 'TNR'
                l_indent = 0
                sp_before = 2
                sp_after = 2
                
            story.append(TOCLineFlowable(entry_title, entry_page, fontName=f_name, fontSize=12, leftIndent=l_indent, spaceBefore=sp_before, spaceAfter=sp_after, usableWidth=usable))
            continue
        # Headings
        if st_name=='Heading 1':
            markup=para_to_markup(para)
            story.append(Paragraph(markup, h1_style))
            continue
        elif st_name=='Heading 2':
            markup=para_to_markup(para)
            story.append(Paragraph(markup, h2_style))
            continue
        elif st_name=='Heading 3':
            markup=para_to_markup(para)
            story.append(Paragraph(markup, h3_style))
            continue
        else:
            # Normal
            # Detect font size to choose caption/source/cover/body
            # Get sizes of non-empty runs
            sizes=[]
            for r in para.runs:
                if r.text.strip():
                    if r.font.size:
                        sizes.append(r.font.size.pt)
            max_sz = max(sizes) if sizes else 12
            # Cover paras 0-5
            if idx<=5:
                bef,aft = cover_bef_aft.get(idx,(0,0))
                # choose 14 vs 12
                if max_sz>=13.5:
                    st = ParagraphStyle(f'Cov14_{idx}', parent=cover14, spaceBefore=bef, spaceAfter=aft)
                else:
                    # cover12 but preserve bold? markup already has bold
                    st = ParagraphStyle(f'Cov12_{idx}', parent=cover12, spaceBefore=bef, spaceAfter=aft)
                markup=para_to_markup(para)
                story.append(Paragraph(markup, st))
                continue
            # Caption 11pt bold? check max_sz 11
            if abs(max_sz-11)<0.2:
                # caption: alignment from docx
                al = para.alignment
                st = caption_center if (al is not None and int(al)==1) else caption_style
                markup=para_to_markup(para)
                story.append(Paragraph(markup, st))
                continue
            if abs(max_sz-10)<0.2:
                al = para.alignment
                st = source_center if (al is not None and int(al)==1) else source_style
                markup=para_to_markup(para)
                story.append(Paragraph(markup, st))
                continue
            # Refs?
            if ref_start is not None and idx>=ref_start:
                markup=para_to_markup(para)
                story.append(Paragraph(markup, ref_style))
                continue
            # Body: check hanging vs first-indent via paragraph_format
            try:
                li = para.paragraph_format.left_indent
                fli = para.paragraph_format.first_line_indent
                has_hang = (li is not None and fli is not None)
            except:
                has_hang=False
            markup=para_to_markup(para)
            # alignment
            al = para.alignment
            if al is not None and int(al)==1:
                story.append(Paragraph(markup, center12))
            elif al is not None and int(al)==0:
                # left aligned? could be caption already handled, else left
                story.append(Paragraph(markup, left12))
            else:
                # justify
                if has_hang:
                    # hanging lists: use body_hang (left 1.25 first -0.4)
                    # But refs already handled; this is enumerations
                    story.append(Paragraph(markup, body_hang))
                else:
                    # check firstLineIndent? body has first 1.25, but some paras like "Sumber data terdiri atas:" have first 1.25? Actually 141 has first 1.25 but short? Keep body_style
                    # For paras like "Ketentuan pengukuran:" etc, also body
                    story.append(Paragraph(markup, body_style))
    elif tag=='tbl':
        # find table
        tbl=None; t_idx=None
        for _i,_t in enumerate(d.tables):
            if _t._tbl is child:
                tbl=_t; t_idx=_i; break
        if tbl is None:
            continue
        nrows=len(tbl.rows); ncols=len(tbl.columns)
        if ncols==5:
            widths=[usable/5]*5
        elif ncols==4:
            widths=[usable/4]*4
        else:
            widths=[usable/ncols]*ncols
        data=[]
        for ri,row in enumerate(tbl.rows):
            rdata=[]
            is_header=(ri==0)
            for cell in row.cells:
                cps=[cp for cp in cell.paragraphs if not (cp.text.strip()=='' and len(cell.paragraphs)>1)]
                if not cps:
                    st = cellH10 if is_header else cell10
                    rdata.append(Paragraph(' ', st))
                elif len(cps)==1:
                    cp=cps[0]
                    markup=para_to_markup(cp)
                    st = cellH10 if is_header else cell10
                    al=cp.alignment
                    if al is not None:
                        v=int(al)
                        if v==1:
                            st=ParagraphStyle('tmpC', parent=st, alignment=TA_CENTER)
                        elif v==3:
                            st=ParagraphStyle('tmpJ', parent=st, alignment=TA_JUSTIFY)
                    rdata.append(Paragraph(markup if markup.strip() else ' ', st))
                else:
                    lst=[]
                    for cp in cps:
                        markup=para_to_markup(cp)
                        st = cellH10 if is_header else cell10
                        lst.append(Paragraph(markup if markup.strip() else ' ', st))
                    rdata.append(lst)
            data.append(rdata)
        from reportlab.lib.colors import HexColor
        hdr_bg = HexColor('#F2F2F2')
        tstyle=TableStyle([
            ('LINEABOVE',(0,0),(-1,0),1,colors.black),
            ('LINEBELOW',(0,0),(-1,0),0.5,colors.black),
            ('LINEBELOW',(0,-1),(-1,-1),1,colors.black),
            ('BACKGROUND',(0,0),(-1,0),hdr_bg),
            ('VALIGN',(0,0),(-1,-1),'TOP'),
            ('LEFTPADDING',(0,0),(-1,-1),4),
            ('RIGHTPADDING',(0,0),(-1,-1),4),
            ('TOPPADDING',(0,0),(-1,-1),3),
            ('BOTTOMPADDING',(0,0),(-1,-1),3),
            ('ALIGN',(0,0),(-1,-1),'CENTER'),
        ])
        rt=Table(data, colWidths=widths, repeatRows=1)
        rt.setStyle(tstyle)
        story.append(rt)
        story.append(Spacer(1,6))

print('story len', len(story), 'sect', sect_idx)
doc_tpl.build(story)
print('BUILT', PDF_PATH, Path(PDF_PATH).stat().st_size)
