import pathlib, re, fitz
from docx import Document
DOCX = pathlib.Path(r"Z:\SKRIPSII\SKRIPSI ARTHUR\skripsi siddharta\01_Naskah_Utama\Proposal_Skripsi_MLBB_GenZ_v2.docx")
PDF = pathlib.Path(r"Z:\SKRIPSII\SKRIPSI ARTHUR\skripsi siddharta\01_Naskah_Utama\Proposal_Skripsi_MLBB_GenZ_v2.pdf")
OUT = pathlib.Path(r"C:\Users\ARTHUR~1\AppData\Local\Temp\opencode\verify_pdf_v2_out.txt")
doc = Document(str(DOCX))
pdf = fitz.open(str(PDF))
def norm(s):
    s=s.replace('\u00a0',' ').replace('\n',' ').replace('\r',' ')
    s=re.sub(r'\s+',' ',s).strip().lower()
    return s
pdf_text = ""
for pg in pdf:
    pdf_text += pg.get_text("text") + "\n"
pdf_norm = norm(pdf_text)
# paras missing
nonempty = [(i,p.text) for i,p in enumerate(doc.paragraphs) if p.text.strip()!='']
missing=[]
for i,txt in nonempty:
    n=norm(txt)
    # TOC entries have tab; pdf has dots between, so check title and page separately?
    if '\t' in txt:
        parts=txt.split('\t')
        title=norm(parts[0])
        page=norm(parts[-1])
        if title not in pdf_norm or page not in pdf_norm:
            # check title words present?
            # For missing, check title substring first 30 chars?
            if title[:30] not in pdf_norm:
                missing.append((i,txt[:80]))
        continue
    # equations have tab + (3.1): check eq part and num separately
    if '(3.1)' in txt or '(3.2)' in txt:
        # split
        pt = txt.split('\t')
        eq = norm(pt[0])
        num = norm(pt[-1])
        if eq not in pdf_norm:
            # try without normalizing subscripts? subscripts preserved, should match
            missing.append((i,txt[:80]+' EQ_MISS'))
        if num not in pdf_norm:
            missing.append((i,txt[:80]+' NUM_MISS'))
        continue
    if n not in pdf_norm:
        # try first 60 chars?
        if n[:60] not in pdf_norm:
            missing.append((i,txt[:100]))
        else:
            # split across lines? check words?
            pass
# cells
cell_missing=[]
for ti,t in enumerate(doc.tables):
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            ct=cell.text.strip()
            if not ct: continue
            n=norm(ct)
            if n not in pdf_norm:
                # try first 50?
                if n[:50] not in pdf_norm:
                    cell_missing.append((ti,ri,ci,ct[:80]))
# keys
keys = ["PROPOSAL TUGAS AKHIR","PENGARUH PEMASARAN INFLUENCER","HALAMAN PERSETUJUAN","Dr Fredella Colline","DAFTAR ISI","DAFTAR TABEL","DAFTAR GAMBAR","BAB I","BAB II","BAB III","1.1 Latar Belakang","2.2 Penelitian Sebelumnya","S01","S18","Tabel 2.1","Tabel 3.1","3.4 Operasionalisasi","DAFTAR PUSTAKA","Ajzen","Winarno","Universitas Kristen Krida Wacana","Ohanian","Colline","(3.1)","(3.2)","Gambar 2.1","Diolah penulis (2026)"]
key_res=[]
for k in keys:
    key_res.append((k, norm(k) in pdf_norm))
# fonts, spans, links, dots
spans=0; italic_spans=0; fonts=set()
for pg in pdf:
    d=pg.get_text("dict")
    for b in d.get("blocks",[]):
        for l in b.get("lines",[]):
            for s in l.get("spans",[]):
                spans+=1
                fonts.add(s.get("font",""))
                if "Italic" in s.get("font","") or "italic" in s.get("font","").lower():
                    italic_spans+=1
links=sum(len(list(pg.annots() or [])) for pg in pdf)
# dots count
dots = pdf_text.count("....")
with OUT.open("w", encoding="utf-8") as f:
    f.write(f"docx paras={len(doc.paragraphs)} nonempty={len(nonempty)} tables={len(doc.tables)} pdf pages={len(pdf)} chars={len(pdf_text)} spans={spans} italic_spans={italic_spans} fonts={fonts} dots4={dots} links_annots={links}\n")
    f.write(f"missing paras {len(missing)}/{len(nonempty)}\n")
    for m in missing[:50]:
        f.write(f"  MISS {m}\n")
    f.write(f"cell missing {len(cell_missing)}\n")
    for m in cell_missing[:50]:
        f.write(f"  CELLMISS {m}\n")
    f.write("keys:\n")
    for k,ok in key_res:
        f.write(f"  {'OK' if ok else 'MISS'} {k}\n")
    # margins? check via construction (4/3/3/3) - report
    f.write(f"pdf size={PDF.stat().st_size} bytes\n")
print("done verify")
