import docx, re
from pathlib import Path
p = Path(r"Z:\SKRIPSII\SKRIPSI ARTHUR\skripsi siddharta\01_Naskah_Utama\Proposal_Skripsi_MLBB_GenZ.docx")
d = docx.Document(str(p))
out = []
out.append(f"paras={len(d.paragraphs)} tables={len(d.tables)} sections={len(d.sections)}")
for i,s in enumerate(d.sections):
    out.append(f"sec{i} L={s.left_margin.cm:.4f} R={s.right_margin.cm:.4f} T={s.top_margin.cm:.4f} B={s.bottom_margin.cm:.4f} hd={s.header_distance.cm:.4f} fd={s.footer_distance.cm:.4f} titlePg={s.different_first_page_header_footer} fmt={s._sectPr.xpath('.//w:pgNumType/@w:fmt')} start={s._sectPr.xpath('.//w:pgNumType/@w:start')}")
# body indent check (sample Normal paras)
bad_indent=[]
for idx,pp in enumerate(d.paragraphs):
    if pp.style.name=='Normal' and pp.text.strip() and len(pp.text)>50:
        # skip cover (center), TOC (toc styles are not Normal), captions (Tabel/Gambar/Sumber), equations (contain 3.1)
        t=pp.text.strip()
        if t.startswith("PROPOSAL") or t.startswith("PENGARUH") or t.startswith("Diajukan") or t.startswith("PROGRAM") or t.startswith("Tabel") or t.startswith("Gambar") or t.startswith("Sumber:") or "(3." in t or t.startswith("Catatan Tabel"):
            continue
        fi = pp.paragraph_format.first_line_indent
        ficm = fi.cm if fi else None
        if ficm is None or abs(ficm-1.25)>0.05:
            bad_indent.append((idx, t[:60], ficm))
        if len(bad_indent)>5: break
out.append(f"body_indent_bad_sample={bad_indent}")
# pustaka hanging
pust_start=None
for idx,pp in enumerate(d.paragraphs):
    if pp.text.strip()=="DAFTAR PUSTAKA" and pp.style.name=='Heading 1':
        pust_start=idx; break
out.append(f"pust_start={pust_start}")
if pust_start:
    refs = d.paragraphs[pust_start+1:]
    out.append(f"refs_count={len(refs)}")
    bad_hang=[]
    for pp in refs[:5]:
        li=pp.paragraph_format.left_indent.cm if pp.paragraph_format.left_indent else None
        fi=pp.paragraph_format.first_line_indent.cm if pp.paragraph_format.first_line_indent else None
        ls=pp.paragraph_format.line_spacing
        out.append(f"ref hang L={li} F={fi} ls={ls} txt={pp.text[:80]}")
    # alphabetical?
    first_words=[pp.text.split(',')[0].strip().lower() for pp in refs]
    sorted_words=sorted(first_words)
    out.append(f"alpha_ok={first_words==sorted_words}")
    # numbered?
    numbered=[pp.text.strip()[:2] for pp in refs if re.match(r'^\d+\.', pp.text.strip())]
    out.append(f"numbered_refs={len(numbered)}")
# cover sizes
for idx in [1,2,3,4,5]:
    pp=d.paragraphs[idx]
    sizes=[(r.text[:30], r.font.size.pt if r.font.size else None, r.bold, r.italic) for r in pp.runs]
    out.append(f"cover p{idx} {pp.text[:60]!r} {sizes} bef={pp.paragraph_format.space_before.pt if pp.paragraph_format.space_before else None} aft={pp.paragraph_format.space_after.pt if pp.paragraph_format.space_after else None}")
# non12 detail
for pp in d.paragraphs:
    for r in pp.runs:
        if not r.text.strip(): continue
        sz=r.font.size.pt if r.font.size else 0
        if abs(sz-12)>0.01:
            # classify
            style=pp.style.name
            txt=pp.text[:40].replace('\n',' ')
            out.append(f"non12 style={style} sz={sz} txt={txt!r} run={r.text[:30]!r}")
            break
    if len([l for l in out if l.startswith('non12')])>30:
        break
# italic et al check
total_et=0; italic_et=0; nonitalic_et=[]
full_runs=[]
for pp in list(d.paragraphs)+[cell.paragraphs[0] for t in d.tables for row in t.rows for cell in row.cells]:
    # need all paras in cells
    pass
# simpler: iterate paras + tables
def iter_all_paras(doc):
    for pp in doc.paragraphs:
        yield pp
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for pp in cell.paragraphs:
                    yield pp
for pp in iter_all_paras(d):
    txt="".join([r.text for r in pp.runs])
    # find et al occurrences with italic?
    # check runs containing et al
    for r in pp.runs:
        if "et al." in r.text:
            total_et+=r.text.count("et al.")
            if r.italic:
                italic_et+=r.text.count("et al.")
            else:
                nonitalic_et.append(pp.text[:80])
out.append(f"et_al total_runs_occ={total_et} italic={italic_et} nonitalic_samples={nonitalic_et[:5]}")
# parenthetical dan vs &
paren_dan=0; paren_amp=0
for pp in iter_all_paras(d):
    for m in re.finditer(r'\([^()]*\d{4}[^()]*\)', pp.text):
        inner=m.group(0)
        if " dan " in inner:
            paren_dan+=1
        if " & " in inner:
            paren_amp+=1
out.append(f"paren_dan={paren_dan} paren_amp={paren_amp}")
# check equations
for pp in d.paragraphs:
    if "(3.1)" in pp.text or "(3.2)" in pp.text:
        tabs=[(t.position.cm if hasattr(t,'position') else 0, str(t.alignment), str(t.leader)) for t in pp.paragraph_format.tab_stops]
        out.append(f"eq {pp.text[:60]!r} tabs={tabs} align={pp.alignment}")
# tables APA?
for ti,t in enumerate(d.tables):
    xml=t._tbl.xml
    has_top=("w:top" in xml); has_bottom=("w:bottom" in xml)
    # check header shading F2F2F2
    has_shd=("F2F2F2" in xml)
    out.append(f"table{ti} {len(t.rows)}x{len(t.columns)} top={has_top} bottom={has_bottom} shd={has_shd}")
# images width
for sh in d.inline_shapes:
    out.append(f"img w_cm={sh.width.cm:.2f} h_cm={sh.height.cm:.2f}")
# footers PAGE field?
for i,s in enumerate(d.sections):
    xml=s.footer._element.xml
    out.append(f"sec{i} footer_has_PAGE={'PAGE' in xml} has_fldChar={'fldChar' in xml}")
# hyperlinks color?
from docx.opc.constants import RELATIONSHIP_TYPE as RT
hlinks=[r for r in d.part.rels.values() if r.reltype==RT.HYPERLINK]
out.append(f"hyperlinks={len(hlinks)}")
# check hyperlink color black via xml search
import zipfile
# count w:color 000000 in document.xml? approximate
xml_all = b"".join([pp._p.xml.encode('utf-8', errors='ignore') for pp in d.paragraphs])
out.append(f"doc_size={p.stat().st_size}")
Path(r"C:\Users\ARTHUR~1\AppData\Local\Temp\opencode\verify_v2_out.txt").write_text("\n".join(out), encoding='utf-8')
# no console print (subscript breaks cp1252)
