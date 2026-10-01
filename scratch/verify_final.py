import pathlib, re, fitz
from docx import Document
DOCX = pathlib.Path(r"Z:\SKRIPSII\SKRIPSI ARTHUR\skripsi siddharta\01_Naskah_Utama\Proposal_Skripsi_MLBB_GenZ_v2.docx")
PDF = pathlib.Path(r"Z:\SKRIPSII\SKRIPSI ARTHUR\skripsi siddharta\01_Naskah_Utama\Proposal_Skripsi_MLBB_GenZ_v2.pdf")
OUT = pathlib.Path(r"C:\Users\ARTHUR~1\AppData\Local\Temp\opencode\verify_final.txt")
doc = Document(str(DOCX))
pdf = fitz.open(str(PDF))
def norm(s):
    s=s.replace('\u00a0',' ').replace('\n',' ').replace('\r',' ')
    s=re.sub(r'\s+',' ',s).strip().lower()
    return s
pdf_text = "".join([pg.get_text("text") for pg in pdf])
pdf_norm = norm(pdf_text)
# headings
from collections import Counter
c=Counter([p.style.name for p in doc.paragraphs])
# italic runs docx
italic_runs = sum(1 for p in doc.paragraphs for r in p.runs if r.italic and r.text.strip())
# spans pdf
spans=0; italic_spans=0; fonts=set(); sizes=set()
for pg in pdf:
    d=pg.get_text("dict")
    for b in d.get("blocks",[]):
        for l in b.get("lines",[]):
            for s in l.get("spans",[]):
                spans+=1
                fonts.add(s.get("font",""))
                sizes.add(round(s.get("size",0),1))
                if "Italic" in s.get("font","") or "ital" in s.get("font","").lower():
                    italic_spans+=1
# links via get_links
total_links=sum(len(pg.get_links()) for pg in pdf)
# dots
dots4=pdf_text.count("....")
# TOC entries count docx toc paras
toc_n=sum(1 for p in doc.paragraphs if p.style.name.startswith('toc'))
# S codes
s_codes=[f"S{i:02d}" for i in range(1,19)]
s_found={s:(s.lower() in pdf_norm) for s in s_codes}
# refs
ref_start=None
for i,p in enumerate(doc.paragraphs):
    if p.text.startswith('Ajzen, I. (1991)'):
        ref_start=i; break
refs = len(doc.paragraphs)-ref_start if ref_start else 0
# equations subscripts in pdf?
has_sub0 = chr(0x2080) in pdf_text
has_sub1 = chr(0x2081) in pdf_text
# footer checks
has_ukrida = "universitas kristen krida wacana" in pdf_norm
# roman footers? search for "| ii", "| iii", "| iv", "| v" and "| 1"
roman_checks = {r:(f"| {r}" in pdf_norm) for r in ["ii","iii","iv","v"]}
arabic_checks = {a:(f"| {a}" in pdf_norm) for a in ["1","2"]}
# gambar
has_gambar_caption = "gambar 2.1" in pdf_norm
has_sumber2026 = "diolah penulis (2026)" in pdf_norm
# image present? check images in pdf
imgs = sum(len(pg.get_images()) for pg in pdf)
with OUT.open("w", encoding="utf-8") as f:
    f.write(f"paras docx={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)} headings={dict(c)}\n")
    f.write(f"italic_runs docx={italic_runs}\n")
    f.write(f"pdf pages={len(pdf)} chars={len(pdf_text)} spans={spans} italic_spans={italic_spans} fonts={fonts} sizes={sorted(sizes)} dots4={dots4} links={total_links} imgs={imgs}\n")
    f.write(f"toc paras docx={toc_n}\n")
    f.write(f"S codes found {sum(s_found.values())}/18 {s_found}\n")
    f.write(f"refs docx from {ref_start} count {refs}\n")
    f.write(f"subscripts 2080={has_sub0} 2081={has_sub1}\n")
    f.write(f"ukrida={has_ukrida} roman={roman_checks} arabic={arabic_checks}\n")
    f.write(f"gambar caption={has_gambar_caption} sumber2026={has_sumber2026}\n")
    # margins construction
    f.write("margins construction 4/3/3/3 A4 usable 396.85pt\n")
    # H1 check
    for i,p in enumerate(doc.paragraphs):
        if p.style.name=='Heading 1':
            f.write(f"H1 {i} :: {repr(p.text[:60])} center={p.alignment} upper={p.text.upper()==p.text}\n")
print("done final")
